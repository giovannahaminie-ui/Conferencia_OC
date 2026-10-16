"""
relatorio_oc.py
─────────────────────────────────────────────────────────────────────────────
Relatório Excel — NF x Ordem de Compra gerada, filtrado por período.

Entram SOMENTE as OCs geradas pelo fluxo, aprovadas, e cuja NF chegou na
coluna "Aguardando Integração" DENTRO do período escolhido (o fluxo só
avança pra Integração depois que a OC é aprovada — SITAPR='APR').

Colunas: Nº NF · Empresa · Filial · Fornecedor · Nº da Ordem de Compra ·
         Situação da OC (E420OCP.SITOCP, lido do Sapiens)

FONTE: banco local (nf_processamento).
  • O período é filtrado pelo criado_em da linha em AGUARDANDO_INTEGRACAO
    (quando a NF entrou nessa coluna).
  • O número da OC vem de divergencias[0].referencia da última linha em
    AGUARDANDO_APROVACAO_OC:  {"apr": "<NUMOCP>", "pen": "<NUMOCP>"}
        apr -> OC consolidada (Juntar_OC)
        pen -> OC nova pendente que foi aprovada (OC_Pendente)

Uso:
    from relatorio_oc import gerar_relatorio_oc_aprovadas
    # sem datas => somente o dia de hoje
    xlsx = gerar_relatorio_oc_aprovadas()
    # período (strings 'YYYY-MM-DD', inclusive nas duas pontas)
    xlsx = gerar_relatorio_oc_aprovadas("2026-08-01", "2026-08-31")
"""

import io
import json
from datetime import date

from db_local import get_local_connection

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError as e:  # pragma: no cover
    raise ImportError(
        "O relatório de OC precisa da biblioteca 'openpyxl'. Instale com:\n"
        r"    .venv\Scripts\python.exe -m pip install openpyxl"
    ) from e


# ── estilos ────────────────────────────────────────────────────────────────
_AZUL_ESCURO = "1F3864"
_AZUL_MEDIO  = "2E5496"
_ZEBRA       = "F2F5FA"

_F_TITULO = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
_F_HEAD   = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
_F_CELL   = Font(name="Calibri", size=11)
_FILL_TIT  = PatternFill("solid", fgColor=_AZUL_ESCURO)
_FILL_HEAD = PatternFill("solid", fgColor=_AZUL_MEDIO)
_FILL_ZEB  = PatternFill("solid", fgColor=_ZEBRA)
_LADO = Side(style="thin", color="BFBFBF")
_BORDA = Border(left=_LADO, right=_LADO, top=_LADO, bottom=_LADO)
_CENTRO = Alignment(horizontal="center", vertical="center")
_ESQ    = Alignment(horizontal="left",  vertical="center")


def _br(iso: str) -> str:
    """'2026-08-31' -> '31/08/2026'"""
    try:
        a, m, d = iso.split("-")
        return f"{d}/{m}/{a}"
    except Exception:
        return iso


def _coletar_linhas(desde=None, ate=None):
    """Retorna (linhas, desde, ate).
    linhas = [(nf, emp, fil, fornecedor, oc, situacao_oc), ...]"""
    hoje = date.today().isoformat()
    desde = (desde or hoje)[:10]
    ate   = (ate   or hoje)[:10]

    conn = get_local_connection()
    cur = conn.cursor()

    # NFs que entraram em AGUARDANDO_INTEGRACAO dentro do período.
    # Com MAX(criado_em) o SQLite traz codemp/codfil/nomfor da linha mais
    # recente (regra de "bare columns" com um único min/max).
    nfs = cur.execute("""
        SELECT numnfc, codemp, codfil, nomfor, MAX(criado_em) AS quando
        FROM nf_processamento
        WHERE status = 'AGUARDANDO_INTEGRACAO'
          AND substr(criado_em, 1, 10) BETWEEN ? AND ?
        GROUP BY numnfc
    """, (desde, ate)).fetchall()

    linhas = []
    for numnfc, codemp, codfil, nomfor, _quando in nfs:
        r = cur.execute("""
            SELECT divergencias FROM nf_processamento
            WHERE numnfc = ? AND status = 'AGUARDANDO_APROVACAO_OC'
            ORDER BY id DESC LIMIT 1
        """, (numnfc,)).fetchone()
        if not r:
            continue  # NF foi direto pra integração — nenhuma OC gerada

        try:
            ref = (json.loads(r[0] or "[]") or [{}])[0].get("referencia") or {}
        except Exception:
            ref = {}
        if not isinstance(ref, dict):
            continue

        for chave in ("apr", "pen"):            # só as OCs GERADAS pelo webservice
            numocp = ref.get(chave)
            if numocp:
                linhas.append((
                    str(numnfc), str(codemp or ""), str(codfil or ""),
                    (nomfor or "").strip(), str(numocp),
                ))

    conn.close()

    # Situação de cada OC — E420OCP.SITOCP lido do Sapiens (Oracle). Se o
    # Oracle estiver indisponível, a coluna sai em branco em vez de derrubar
    # o relatório inteiro.
    situ = {}
    try:
        from oc_service import buscar_situacao_ocs
        situ = buscar_situacao_ocs((l[1], l[2], l[4]) for l in linhas)
    except Exception as e:  # pragma: no cover
        print(f"[relatorio_oc] SITOCP indisponível: {e!r}")

    linhas = [l + (situ.get((l[1], l[2], l[4]), ""),) for l in linhas]

    def _ordkey(l):
        try:
            return (0, int(l[0]))
        except ValueError:
            return (1, l[0])

    return sorted(set(linhas), key=_ordkey), desde, ate


def gerar_relatorio_oc_aprovadas(desde=None, ate=None) -> bytes:
    linhas, desde, ate = _coletar_linhas(desde, ate)

    wb = Workbook()
    ws = wb.active
    ws.title = "OCs Geradas"

    colunas = [
        ("Nº NF",                 16, _CENTRO),
        ("Empresa",               12, _CENTRO),
        ("Filial",                10, _CENTRO),
        ("Fornecedor",            46, _ESQ),
        ("Nº da Ordem de Compra", 24, _CENTRO),
        ("Situação da OC",         22, _ESQ),
    ]
    ncol = len(colunas)

    periodo = _br(desde) if desde == ate else f"{_br(desde)} a {_br(ate)}"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    c = ws.cell(row=1, column=1, value=f"OCs geradas e aprovadas — {periodo}")
    c.font, c.fill, c.alignment = _F_TITULO, _FILL_TIT, _ESQ
    ws.row_dimensions[1].height = 24

    for j, (nome, larg, _al) in enumerate(colunas, start=1):
        cel = ws.cell(row=2, column=j, value=nome)
        cel.font, cel.fill, cel.alignment, cel.border = _F_HEAD, _FILL_HEAD, _CENTRO, _BORDA
        ws.column_dimensions[get_column_letter(j)].width = larg
    ws.row_dimensions[2].height = 22

    for i, linha in enumerate(linhas, start=3):
        for j, (_nome, _larg, alinh) in enumerate(colunas, start=1):
            cel = ws.cell(row=i, column=j, value=linha[j - 1])
            cel.font, cel.border, cel.alignment = _F_CELL, _BORDA, alinh
            if i % 2 == 1:
                cel.fill = _FILL_ZEB

    ws.freeze_panes = "A3"
    if linhas:
        ws.auto_filter.ref = f"A2:{get_column_letter(ncol)}{len(linhas) + 2}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    with open("relatorio_oc_teste.xlsx", "wb") as f:
        f.write(gerar_relatorio_oc_aprovadas("2000-01-01", "2100-01-01"))
    print("relatorio_oc_teste.xlsx gerado.")